"""Seed departments, roles, and employees into the database from 01_LedgerForge_Python_Master_Dataset.xlsx.

Usage (from repo root):
    python scripts/seed_from_dataset.py [path_to_excel]

Defaults to '01_LedgerForge_Python_Master_Dataset.xlsx' if no argument provided.
"""

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

# Add project root to sys.path so 'app' imports work
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.db.session import get_engine  # noqa: E402
from app.models import Department, Employee, Role  # noqa: E402


def _to_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def seed(xlsx_path: Path):
    if not xlsx_path.exists():
        print(f"Error: dataset file not found at {xlsx_path}")
        sys.exit(1)

    print(f"Loading dataset from: {xlsx_path.name}...")
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)

    engine = get_engine()

    with Session(engine) as session:
        # 1. Seed Departments
        dept_sheet = wb["departments"]
        dept_rows = list(dept_sheet.iter_rows(values_only=True))
        dept_header = dept_rows[0]
        dept_count = 0

        existing_depts = set(session.scalars(select(Department.department_id)).all())

        for row in dept_rows[1:]:
            rec = dict(zip(dept_header, row))
            if not rec.get("department_id") or rec["department_id"] in existing_depts:
                continue
            dept = Department(
                department_id=str(rec["department_id"]).strip(),
                department_name=str(rec["department_name"]).strip(),
                cost_center=str(rec["cost_center"]).strip(),
                description=str(rec["description"]).strip() if rec.get("description") else None,
            )
            session.add(dept)
            existing_depts.add(dept.department_id)
            dept_count += 1

        session.commit()
        print(f"-> Seeded {dept_count} new departments (total: {len(existing_depts)})")

        # 2. Seed Roles
        role_sheet = wb["roles"]
        role_rows = list(role_sheet.iter_rows(values_only=True))
        role_header = role_rows[0]
        role_count = 0

        existing_roles = set(session.scalars(select(Role.role_id)).all())

        for row in role_rows[1:]:
            rec = dict(zip(role_header, row))
            if not rec.get("role_id") or rec["role_id"] in existing_roles:
                continue
            role = Role(
                role_id=str(rec["role_id"]).strip(),
                department_id=str(rec["department_id"]).strip(),
                role_title=str(rec["role_title"]).strip(),
                level=int(rec["level"]),
                grade=str(rec["grade"]).strip(),
                band_min_annual=float(rec["band_min_annual"]),
                band_mid_annual=float(rec["band_mid_annual"]),
                band_max_annual=float(rec["band_max_annual"]),
                overtime_eligible=(str(rec.get("overtime_eligible", "")).upper() == "Y"),
            )
            session.add(role)
            existing_roles.add(role.role_id)
            role_count += 1

        session.commit()
        print(f"-> Seeded {role_count} new roles (total: {len(existing_roles)})")

        # 3. Seed Employees
        emp_sheet = wb["employees"]
        emp_rows = list(emp_sheet.iter_rows(values_only=True))
        emp_header = emp_rows[0]
        emp_count = 0

        existing_emps = set(session.scalars(select(Employee.employee_id)).all())

        for row in emp_rows[1:]:
            rec = dict(zip(emp_header, row))
            emp_id = rec.get("employee_id")
            if not emp_id or emp_id in existing_emps:
                continue

            manager_id = rec.get("manager_id")
            if manager_id:
                manager_id = str(manager_id).strip()

            emp = Employee(
                employee_id=str(emp_id).strip(),
                first_name=str(rec["first_name"]).strip(),
                last_name=str(rec["last_name"]).strip(),
                work_email=str(rec["work_email"]).strip(),
                gender=str(rec["gender"]).strip(),
                date_of_birth=_to_date(rec["date_of_birth"]),
                hire_date=_to_date(rec["hire_date"]),
                exit_date=_to_date(rec.get("exit_date")),
                employment_status=str(rec["employment_status"]).strip(),
                employment_type=str(rec["employment_type"]).strip(),
                department_id=str(rec["department_id"]).strip(),
                role_id=str(rec["role_id"]).strip(),
                manager_id=manager_id,
                work_location=str(rec["work_location"]).strip(),
                work_mode=str(rec["work_mode"]).strip(),
            )
            session.add(emp)
            existing_emps.add(emp.employee_id)
            emp_count += 1

        session.commit()
        print(f"-> Seeded {emp_count} new employees (total: {len(existing_emps)})")

    print("Seeding completed successfully!")


if __name__ == "__main__":
    file_path = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else ROOT / "01_LedgerForge_Python_Master_Dataset.xlsx"
    )
    seed(file_path)
